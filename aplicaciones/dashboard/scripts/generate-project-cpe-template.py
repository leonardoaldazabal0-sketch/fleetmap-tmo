import argparse
from pathlib import Path
from shutil import copyfile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation


DEFAULT_OUTPUT = Path(r"F:\archivos\propuesta\Plantilla_Solicitud_Cliente_CPE.xlsx")
DEFAULT_DEPLOY_OUTPUT = Path(r"F:\CPE-TMO\documentacion\Plantilla_Solicitud_Cliente_CPE.xlsx")
MAX_RADIOS = 300
DATA_START = 6
DATA_END = DATA_START + MAX_RADIOS - 1

parser = argparse.ArgumentParser(description="Genera la plantilla dinámica del Proyecto CPE")
parser.add_argument("--groups", type=int, default=10, help="Cantidad de grupos TMO (1-100)")
parser.add_argument("--dmo-groups", type=int, default=10, help="Cantidad de grupos DMO (0-50)")
parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
parser.add_argument("--no-deploy", action="store_true")
args = parser.parse_args()
if not 1 <= args.groups <= 100:
    parser.error("--groups debe estar entre 1 y 100")
if not 0 <= args.dmo_groups <= 50:
    parser.error("--dmo-groups debe estar entre 0 y 50")

OUTPUT = args.output.resolve()
DEPLOY_OUTPUT = DEFAULT_DEPLOY_OUTPUT
GROUPS = args.groups
DMO_GROUPS = args.dmo_groups

NAVY = "1F4E78"
SLATE = "5B6573"
BLUE = "D9EAF7"
PALE_BLUE = "EEF5FA"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
GRID = "A6A6A6"
RED = "F4CCCC"
INTERNAL_GRAY = "E7E6E6"

thin = Side(style="thin", color=GRID)
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def setup(ws):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35


def style(cell, fill=WHITE, color="000000", bold=False, size=10, align="center"):
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Arial", size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = border


def merge_heading(ws, row, start, end, value, fill=NAVY):
    ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    ws.cell(row, start, value)
    style(ws.cell(row, start), fill, WHITE, True, 11)
    for col in range(start + 1, end + 1):
        ws.cell(row, col).border = border


def add_x_validation(ws, ranges):
    dv = DataValidation(type="list", formula1='"X"', allow_blank=True)
    dv.errorTitle = "Valor no permitido"
    dv.error = "Use X o deje la celda vacía."
    dv.promptTitle = "Asignación"
    dv.prompt = "X = solicitado; vacío = no solicitado."
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    for target in ranges:
        dv.add(target)


def add_integer_validation(ws, ranges):
    for target in ranges:
        min_col, min_row, _, _ = range_boundaries(target)
        anchor = f"{get_column_letter(min_col)}{min_row}"
        dv = DataValidation(
            type="custom",
            formula1=(
                f'=OR({anchor}="",AND(IFERROR(--{anchor},0)>=1,'
                f'IFERROR(--{anchor},0)<=16777215,MOD(IFERROR(--{anchor},0),1)=0))'
            ),
            allow_blank=True,
        )
        dv.errorTitle = "ISSI/GSSI inválido"
        dv.error = "Ingrese un número entero entre 1 y 16777215."
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(target)


def add_label_validation(ws, target, label="Label del grupo", max_length=15):
    dv = DataValidation(
        type="textLength", operator="lessThanOrEqual", formula1=str(max_length), allow_blank=True
    )
    dv.errorTitle = "Label demasiado largo"
    dv.error = f"{label} admite un máximo de {max_length} caracteres."
    dv.promptTitle = label
    dv.prompt = f"Máximo {max_length} caracteres."
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(target)


def add_list_validation(ws, target, defined_name):
    dv = DataValidation(type="list", formula1=f"={defined_name}", allow_blank=True)
    dv.errorTitle = "Valor no permitido"
    dv.error = "Seleccione un valor de la lista."
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(target)


def set_widths(ws, metadata_count, group_start, group_end, last_col):
    for col in range(1, last_col + 1):
        if col == 1:
            width = 6
        elif col == 2:
            width = 16
        elif col in (3, 5, 7):
            width = 15
        elif col == 4:
            width = 25
        elif col == 6:
            width = 18
        elif col == 8:
            width = 18
        elif col == 9:
            width = 20
        elif col in (10, 11):
            width = 18
        elif group_start <= col <= group_end:
            width = 12
        else:
            width = 14
        ws.column_dimensions[get_column_letter(col)].width = width


wb = Workbook()
wb.remove(wb.active)
wb.properties.title = "Plantilla de programación - Cliente"
wb.properties.subject = "Solicitud de programación de radios TETRA"
wb.properties.creator = "Dolphin Telecom - Proyecto CPE"
wb.properties.description = "Formato dinámico CLIENTE_CPE_V7 con agenda, gateway GPS y grupos configurables."
wb.calculation.calcMode = "auto"
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True

# PLANTILLA PRINCIPAL
ws = wb.create_sheet("PLANTILLA")
setup(ws)
metadata = [
    "N°",
    "FECHA SOLICITUD",
    "RUC",
    "RAZÓN SOCIAL",
    "CONTRATO / OC",
    "NOMBRE EN AGENDA\n(OPCIONAL)",
    "ISSI\n(USO INTERNO)",
    "SERIE\n(USO INTERNO)",
    "TEI\n(USO INTERNO)",
    "MODELO SOLICITADO",
    "ÁREA / USUARIO",
]
group_start = len(metadata) + 1
group_end = group_start + GROUPS - 1
features = [
    "EMERGENCIA\nMODO",
    "EMERGENCIA\nGSSI",
    "LLAMADA PRIVADA\n(X)",
    "CORTE DE LLAMADA\n(X)",
    "ESCANEO\n(X)",
    "DMO\n(X)",
    "BROADCAST ESCUCHA\n(X)",
    "BROADCAST\nGSSI",
    "GPS\n(X)",
    "GPS GATEWAY\nISSI",
    "OTAP\n(X)",
    "REPORTE GPS\n(X)",
    "MAN DOWN\n(X)",
    "COBERTURA /\nLUGAR",
]
feature_start = group_end + 1
feature_end = feature_start + len(features) - 1
last_col = feature_end
last_letter = get_column_letter(last_col)

merge_heading(ws, 3, 1, len(metadata), "PLANTILLA DE PROGRAMACIÓN - CLIENTE")
merge_heading(
    ws,
    3,
    group_start,
    group_end,
    "GRUPOS TETRA  |  LABEL (FILA 4) - GSSI (FILA 5)  |  MARCAR X  |  MÁXIMO 15 CARACTERES",
    SLATE,
)
merge_heading(ws, 3, feature_start, feature_end, "FUNCIONES Y SERVICIOS", NAVY)

for col, header in enumerate(metadata, 1):
    ws.merge_cells(start_row=4, start_column=col, end_row=5, end_column=col)
    ws.cell(4, col, header)
    header_fill = INTERNAL_GRAY if col in (7, 8, 9) else BLUE
    style(ws.cell(4, col), header_fill, "000000", True, 10)
    ws.cell(4, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for index in range(GROUPS):
    col = group_start + index
    ws.cell(4, col, "")
    ws.cell(5, col, "")
    style(ws.cell(4, col), YELLOW, "000000", True, 10)
    style(ws.cell(5, col), PALE_BLUE, "000000", True, 10)

for offset, header in enumerate(features):
    col = feature_start + offset
    ws.merge_cells(start_row=4, start_column=col, end_row=5, end_column=col)
    ws.cell(4, col, header)
    style(ws.cell(4, col), BLUE, "000000", True, 9)

for row in range(DATA_START, DATA_END + 1):
    ws.cell(row, 1, f'=IF(COUNTA(B{row}:K{row})=0,"",ROW()-5)')
    for col in range(1, last_col + 1):
        cell = ws.cell(row, col)
        cell_fill = INTERNAL_GRAY if col in (7, 8, 9) else WHITE
        style(cell, cell_fill, "000000", False, 10)
        if group_start <= col <= group_end:
            cell.font = Font(name="Arial Narrow", size=11, bold=True)
    for col in (7, 8, 9, feature_start + 9):
        ws.cell(row, col).number_format = "@"

set_widths(ws, len(metadata), group_start, group_end, last_col)
for col in range(feature_start, feature_end + 1):
    ws.column_dimensions[get_column_letter(col)].width = 14
ws.column_dimensions[get_column_letter(feature_end)].width = 21
ws.row_dimensions[3].height = 24
ws.row_dimensions[4].height = 34
ws.row_dimensions[5].height = 24

emergency_mode = get_column_letter(feature_start)
emergency_gssi = get_column_letter(feature_start + 1)
private_call = get_column_letter(feature_start + 2)
call_cut = get_column_letter(feature_start + 3)
scan = get_column_letter(feature_start + 4)
dmo = get_column_letter(feature_start + 5)
broadcast = get_column_letter(feature_start + 6)
broadcast_gssi = get_column_letter(feature_start + 7)
gps = get_column_letter(feature_start + 8)
gps_gateway = get_column_letter(feature_start + 9)
otap = get_column_letter(feature_start + 10)
gps_report = get_column_letter(feature_start + 11)
man_down = get_column_letter(feature_start + 12)

add_label_validation(ws, f"{get_column_letter(group_start)}4:{get_column_letter(group_end)}4")
add_label_validation(ws, f"F{DATA_START}:F{DATA_END}", "Nombre en agenda", 12)
add_integer_validation(ws, [
    f"{get_column_letter(group_start)}5:{get_column_letter(group_end)}5",
    f"G{DATA_START}:G{DATA_END}",
    f"{emergency_gssi}{DATA_START}:{emergency_gssi}{DATA_END}",
    f"{broadcast_gssi}{DATA_START}:{broadcast_gssi}{DATA_END}",
    f"{gps_gateway}{DATA_START}:{gps_gateway}{DATA_END}",
])
add_list_validation(ws, f"{emergency_mode}{DATA_START}:{emergency_mode}{DATA_END}", "LISTA_EMERGENCIA")
add_list_validation(ws, f"J{DATA_START}:J{DATA_END}", "LISTA_MODELOS")
add_x_validation(ws, [
    f"{get_column_letter(group_start)}{DATA_START}:{get_column_letter(group_end)}{DATA_END}",
    f"{private_call}{DATA_START}:{private_call}{DATA_END}",
    f"{call_cut}{DATA_START}:{call_cut}{DATA_END}",
    f"{scan}{DATA_START}:{scan}{DATA_END}",
    f"{dmo}{DATA_START}:{dmo}{DATA_END}",
    f"{broadcast}{DATA_START}:{broadcast}{DATA_END}",
    f"{gps}{DATA_START}:{gps}{DATA_END}",
    f"{otap}{DATA_START}:{otap}{DATA_END}",
    f"{gps_report}{DATA_START}:{gps_report}{DATA_END}",
    f"{man_down}{DATA_START}:{man_down}{DATA_END}",
])
ws.conditional_formatting.add(
    f"{emergency_gssi}{DATA_START}:{emergency_gssi}{DATA_END}",
    FormulaRule(
        formula=[f'AND({emergency_mode}{DATA_START}="NO TÁCTICO",{emergency_gssi}{DATA_START}="")'],
        fill=PatternFill("solid", fgColor=RED),
    ),
)
for identity_col in ("G", "H", "I"):
    ws.conditional_formatting.add(
        f"{identity_col}{DATA_START}:{identity_col}{DATA_END}",
        FormulaRule(
            formula=[
                f'AND({identity_col}{DATA_START}<>"",COUNTIF(${identity_col}${DATA_START}:${identity_col}${DATA_END},{identity_col}{DATA_START})>1)'
            ],
            fill=PatternFill("solid", fgColor=RED),
        ),
    )
ws.conditional_formatting.add(
    f"{broadcast_gssi}{DATA_START}:{broadcast_gssi}{DATA_END}",
    FormulaRule(
        formula=[f'AND({broadcast}{DATA_START}="X",{broadcast_gssi}{DATA_START}="")'],
        fill=PatternFill("solid", fgColor=RED),
    ),
)
ws.conditional_formatting.add(
    f"{gps_gateway}{DATA_START}:{gps_gateway}{DATA_END}",
    FormulaRule(
        formula=[f'AND({gps}{DATA_START}="X",{gps_gateway}{DATA_START}="")'],
        fill=PatternFill("solid", fgColor=RED),
    ),
)
ws.conditional_formatting.add(
    f"F{DATA_START}:F{DATA_END}",
    FormulaRule(
        formula=[f'AND(F{DATA_START}<>"",COUNTIF($F${DATA_START}:$F${DATA_END},F{DATA_START})>1)'],
        fill=PatternFill("solid", fgColor=RED),
    ),
)
ws.freeze_panes = "A6"
ws.print_title_rows = "3:5"
ws.print_area = f"A3:{last_letter}{DATA_END}"


def assignment_sheet(name, dmo_sheet=False):
    sheet = wb.create_sheet(name)
    setup(sheet)
    meta = ["N°", "CONTACTO", "ISSI", "ÁREA / USUARIO"]
    group_count = DMO_GROUPS if dmo_sheet else GROUPS
    matrix_start = len(meta) + 1
    matrix_end = matrix_start + group_count - 1
    merge_heading(sheet, 3, 1, len(meta), f"ASIGNACIÓN DE GRUPOS - {name}")
    message = (
        "GRUPOS DMO  |  LABEL (FILA 4) - GSSI (FILA 5)  |  MÁXIMO 15 CARACTERES"
        if dmo_sheet
        else "GRUPOS TETRA DEL FLEET MAP - MARCAR X"
    )
    if group_count:
        merge_heading(sheet, 3, matrix_start, matrix_end, message, SLATE)

    for col, header in enumerate(meta, 1):
        sheet.merge_cells(start_row=4, start_column=col, end_row=5, end_column=col)
        sheet.cell(4, col, header)
        style(sheet.cell(4, col), BLUE, "000000", True, 10)

    for index in range(group_count):
        col = matrix_start + index
        main_col = get_column_letter(group_start + index)
        if dmo_sheet:
            sheet.cell(4, col, "")
            sheet.cell(5, col, "")
            top_fill = YELLOW
            bottom_fill = PALE_BLUE
        else:
            sheet.cell(4, col, f"='PLANTILLA'!{main_col}4")
            sheet.cell(5, col, f"='PLANTILLA'!{main_col}5")
            top_fill = BLUE
            bottom_fill = PALE_BLUE
        style(sheet.cell(4, col), top_fill, "000000", True, 10)
        style(sheet.cell(5, col), bottom_fill, "000000", True, 10)

    for row in range(DATA_START, DATA_END + 1):
        sheet.cell(row, 1, f'=IF(\'PLANTILLA\'!G{row}="","",ROW()-5)')
        sheet.cell(row, 2, f'=IF(\'PLANTILLA\'!F{row}="","",\'PLANTILLA\'!F{row})')
        sheet.cell(row, 3, f'=IF(\'PLANTILLA\'!G{row}="","",\'PLANTILLA\'!G{row})')
        sheet.cell(row, 4, f'=IF(\'PLANTILLA\'!K{row}="","",\'PLANTILLA\'!K{row})')
        for col in range(1, matrix_end + 1):
            style(sheet.cell(row, col), WHITE, "000000", col > len(meta), 10)

    for col in range(1, matrix_end + 1):
        if col == 1:
            width = 6
        elif col == 2:
            width = 20
        elif col == 3:
            width = 15
        elif col == 4:
            width = 20
        else:
            width = 12
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.row_dimensions[3].height = 24
    sheet.row_dimensions[4].height = 34
    sheet.row_dimensions[5].height = 24
    if group_count:
        add_x_validation(sheet, [
            f"{get_column_letter(matrix_start)}{DATA_START}:{get_column_letter(matrix_end)}{DATA_END}"
        ])
    if dmo_sheet and group_count:
        add_label_validation(sheet, f"{get_column_letter(matrix_start)}4:{get_column_letter(matrix_end)}4")
        add_integer_validation(sheet, [
            f"{get_column_letter(matrix_start)}5:{get_column_letter(matrix_end)}5"
        ])
    sheet.freeze_panes = "A6"
    sheet.print_title_rows = "3:5"
    sheet.print_area = f"A3:{get_column_letter(matrix_end)}{DATA_END}"
    return sheet


assignment_sheet("ESCANEO")
assignment_sheet("DMO", dmo_sheet=True)

# Hoja oculta para listas desplegables y reconocimiento del formato.
ws = wb.create_sheet("OP CAMPOS")
emergency_values = ["TÁCTICO", "NO TÁCTICO", "INACTIVO"]
models = ["MTP3550", "MTP3250", "MTP850S", "MTP8550EX", "MTM5400", "OTRO"]
for row, value in enumerate(emergency_values, 2):
    ws.cell(row, 1, value)
for row, value in enumerate(models, 2):
    ws.cell(row, 2, value)
ws["D1"] = "SCHEMA"
ws["D2"] = "CLIENTE_CPE_V7"
ws["E1"] = "DATA_START"
ws["E2"] = DATA_START
ws["F1"] = "MAX_RADIOS"
ws["F2"] = MAX_RADIOS
ws["G1"] = "GROUPS"
ws["G2"] = GROUPS
ws["H1"] = "DMO_GROUPS"
ws["H2"] = DMO_GROUPS
ws.sheet_state = "hidden"

wb.defined_names.add(DefinedName("LISTA_EMERGENCIA", attr_text="'OP CAMPOS'!$A$2:$A$4"))
wb.defined_names.add(DefinedName("LISTA_MODELOS", attr_text="'OP CAMPOS'!$B$2:$B$7"))
wb.active = 0
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)

# Validación final del archivo entregable.
check = load_workbook(OUTPUT, data_only=False)
if check.sheetnames != ["PLANTILLA", "ESCANEO", "DMO", "OP CAMPOS"]:
    raise RuntimeError("El libro generado no contiene las hojas esperadas")
if check["OP CAMPOS"]["D2"].value != "CLIENTE_CPE_V7":
    raise RuntimeError("No se conservó la versión del formato")
if check["OP CAMPOS"].sheet_state != "hidden":
    raise RuntimeError("La hoja técnica debe estar oculta")
if check["PLANTILLA"].max_column != 25 + GROUPS:
    raise RuntimeError("La plantilla no conserva la cantidad solicitada de grupos TMO")
if check["ESCANEO"].max_column != 4 + GROUPS:
    raise RuntimeError("Escaneo no replica la cantidad solicitada de grupos TMO")
if check["DMO"].max_column != max(4, 4 + DMO_GROUPS):
    raise RuntimeError("DMO no conserva la cantidad solicitada de grupos")
if check["PLANTILLA"].max_row < DATA_END:
    raise RuntimeError("La plantilla no conserva las filas disponibles")
check.close()

if not args.no_deploy:
    DEPLOY_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    copyfile(OUTPUT, DEPLOY_OUTPUT)

print(OUTPUT)
