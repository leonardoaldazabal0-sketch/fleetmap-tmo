import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SOURCE = Path(sys.argv[1])
OUTPUT = Path(sys.argv[2])
data = json.loads(SOURCE.read_text(encoding="utf-8"))

NAVY = "1F4E78"
BLUE = "D9EAF7"
PALE = "F4F7F8"
WHITE = "FFFFFF"
GREEN = "DFF2E8"
YELLOW = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
BORDER_COLOR = "B7C4C8"
thin = Side(style="thin", color=BORDER_COLOR)
border = Border(left=thin, right=thin, top=thin, bottom=thin)

capabilities = ["Grupos TMO", "Emergencia", "Escaneo", "DMO", "Broadcast", "GPS", "OTAP", "Agenda"]


def setup(ws):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def header(ws, row=1):
    for cell in ws[row]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[row].height = 32


def body(ws):
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border


def widths(ws, values):
    for index, width in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def color_status_cells(ws, columns, start=2):
    for col in columns:
        letter = get_column_letter(col)
        area = f"{letter}{start}:{letter}{ws.max_row}"
        ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=['"SOPORTADO"'], fill=PatternFill("solid", fgColor=GREEN)))
        ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=['"PARCIAL"'], fill=PatternFill("solid", fgColor=YELLOW)))
        ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=['"NO SOPORTADO"'], fill=PatternFill("solid", fgColor=RED)))
        ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=['"SIN CPE VÁLIDO"'], fill=PatternFill("solid", fgColor=GRAY)))


wb = Workbook()
summary = wb.active
summary.title = "Resumen"
summary.sheet_view.showGridLines = False
summary.merge_cells("A1:H1")
summary["A1"] = "AUDITORÍA DE COBERTURA DEL PARSER CPE"
summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
summary["A1"].font = Font(name="Arial", size=14, bold=True, color=WHITE)
summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
summary.row_dimensions[1].height = 28
summary.append(["Generado", data["generatedAt"], "Familias CP", data["summary"]["namedFamilies"], "Variantes familia/versión", data["summary"]["familyVersionVariants"], "CPE válidos", data["summary"]["validFiles"]])
summary.append(["Base SQLite", data["databasePath"], "CPE con error", data["summary"]["errorFiles"], "Variantes sin CP", data["summary"]["unknownFamilyVariants"], "Criterio", "Activo/Inactivo o estructura validada"])
summary.append([])
summary.append(["Función", "Variantes soportadas", "Parciales", "No soportadas", "Sin CPE válido", "Familias con brecha", "CPE legibles", "Cobertura %"])
for item in data["capabilities"]:
    summary.append([item["name"], item["SOPORTADO"], item["PARCIAL"], item["NO SOPORTADO"], item["SIN CPE VÁLIDO"], item["familiesWithGaps"], item["readableFiles"], item["coveragePercent"]])
for cell in summary[5]:
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(name="Arial", size=10, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
for row in summary.iter_rows(min_row=2):
    for cell in row:
        if cell.row != 5:
            cell.font = Font(name="Arial", size=9, bold=cell.column % 2 == 1 and cell.row in (2, 3))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
summary.freeze_panes = "A5"
summary.auto_filter.ref = f"A5:H{summary.max_row}"
widths(summary, [20, 24, 18, 18, 18, 19, 16, 14])


def coverage_rows(source_rows):
    result = []
    for item in source_rows:
        pending = ", ".join(item["pending"])
        impact = "ALTO" if item["pending"] and item["validFiles"] >= 500 else "MEDIO" if item["pending"] and item["validFiles"] >= 50 else "BAJO" if item["pending"] else "CUBIERTO"
        result.append([
            impact, item["codeplug"], item["software"], item["cpeVersion"], item["cpsLength"] or "No determinado",
            item["validFiles"], item["errorFiles"], item["exampleIssi"],
            *[item["coverage"][name] for name in capabilities],
            pending, item["requiredFiles"], item["examplePath"],
        ])
    return result


coverage_headers = ["Impacto", "Familia CP", "Versión SW", "Versión CPE", "Tamaño CPS", "CPE válidos", "CPE con error", "ISSI ejemplo", *capabilities, "Pendientes", "Archivos requeridos", "Ruta de ejemplo"]
coverage = wb.create_sheet("Cobertura por versión")
coverage.append(coverage_headers)
for row in coverage_rows(data["rows"]):
    coverage.append(row)
setup(coverage)
header(coverage)
body(coverage)
widths(coverage, [11, 12, 19, 12, 15, 12, 13, 13, 15, 15, 15, 15, 15, 15, 15, 15, 30, 54, 70])
color_status_cells(coverage, range(9, 17))

pending_rows = [row for row in data["rows"] if row["pending"]]
pending = wb.create_sheet("Pendientes")
pending.append(coverage_headers)
for row in coverage_rows(pending_rows):
    pending.append(row)
setup(pending)
header(pending)
body(pending)
widths(pending, [11, 12, 19, 12, 15, 12, 13, 13, 15, 15, 15, 15, 15, 15, 15, 15, 30, 54, 70])
color_status_cells(pending, range(9, 17))

families = wb.create_sheet("Resumen por familia")
families.append(["Familia CP", "Variantes", "CPE válidos", "CPE con error", *capabilities, "Funciones con brecha"])
for item in data["families"]:
    families.append([item["codeplug"], item["variants"], item["validFiles"], item["errorFiles"], *[item["coverage"][name] for name in capabilities], ", ".join(item["pending"])])
setup(families)
header(families)
body(families)
widths(families, [12, 11, 13, 13, 15, 15, 15, 15, 15, 15, 15, 15, 38])
color_status_cells(families, range(5, 13))

legend = wb.create_sheet("Leyenda")
legend.append(["Estado", "Significado"])
legend.append(["SOPORTADO", data["definitions"]["supported"]])
legend.append(["PARCIAL", data["definitions"]["partial"]])
legend.append(["NO SOPORTADO", data["definitions"]["unsupported"]])
legend.append(["SIN CPE VÁLIDO", "Solo hay archivos con error para esa combinación; no se evalúa cobertura."])
legend.append([])
legend.append(["Función", "Qué se considera legible"])
legend.append(["Grupos TMO", "Existe una estructura de grupos validada para la familia y tamaño CPS."])
legend.append(["Emergencia", "El parser reconoce modo táctico/no táctico y su destino."])
legend.append(["Escaneo / DMO / GPS / OTAP", "La función devuelve exclusivamente Activo o Inactivo; los valores no verificados no cuentan como soporte."])
legend.append(["Broadcast", "Devuelve Inactivo, Solo escucha o Revisar con estructura validada."])
legend.append(["Agenda", "Existe una estructura de contactos validada contra User Report para ese tamaño CPS."])
setup(legend)
header(legend)
body(legend)
widths(legend, [28, 110])

for ws in wb.worksheets:
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.orientation = "landscape"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)
print(OUTPUT)
